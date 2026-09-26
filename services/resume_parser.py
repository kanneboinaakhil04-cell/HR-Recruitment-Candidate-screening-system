import os
import re
import pypdf
import docx

def extract_text_from_pdf(file_path):
    """
    Extract text content from PDF file using pypdf.
    Handles multipage documents, formatted text, and unexpected stream errors.
    """
    text = ""
    try:
        reader = pypdf.PdfReader(file_path)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        print(f"[ResumeParser Error] Failed to parse PDF {file_path}: {e}")
    return text.strip()


def extract_text_from_docx(file_path):
    """
    Extract text content from Word DOCX file using python-docx.
    Extracts paragraphs, tables, and bullet points.
    """
    text = ""
    try:
        doc = docx.Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text += paragraph.text + "\n"
        
        # Extract text inside table cells if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"[ResumeParser Error] Failed to parse DOCX {file_path}: {e}")
    return text.strip()


def extract_resume_text(file_path):
    """
    Dispatches file to appropriate extractor based on file extension.
    Validates file existence and unsupported extensions.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found at path: {file_path}")
    
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        text = extract_text_from_pdf(file_path)
    elif ext == '.docx':
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX files are supported.")
    
    # Basic text normalization (collapsing multiple whitespace)
    cleaned_text = re.sub(r'\s+', ' ', text)
    return cleaned_text
